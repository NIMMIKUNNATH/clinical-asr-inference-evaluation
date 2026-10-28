"""Five-member panel (180 Fareez events) with the physicians' UPDATED ratings, computed exactly as in
five_panel_agreement.py / plot_fig7_five_reviewer.py (Fleiss complete case, bootstrap 2,000, seed 20260915,
majority with ties), and the updated Figure 7.

Students: clinical_review_inputs/student_<n>_180.xlsx (sheet 'ASR Comparison - 180').
Physicians: 180-event rows of physician_1_312.xlsx / physician_2_312.xlsx, matched on Event ID.
"""
import sys
from collections import Counter
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib import font_manager  # noqa: E402

D = Path("clinical_review_inputs")
U = Path(sys.argv[1])
OUT = Path(__file__).parent
LAB = ["No", "Partially", "Yes"]
INK, FIRE, GREY = "#2D8098", "#D15E5C", "#80909D"


def norm(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    t = str(v).strip().lower()
    if t.startswith("partial"):
        return "Partially"
    return {"yes": "Yes", "y": "Yes", "no": "No", "n": "No"}.get(t)


def read_student(p):
    wb = openpyxl.load_workbook(D / p, data_only=True, read_only=True)
    ws = wb["ASR Comparison - 180"]
    ids, meta, mean, flag = [], {}, {}, {}
    for r in ws.iter_rows(min_row=2, max_col=9, values_only=True):
        if r[0] is None:
            continue
        i = str(r[0]).strip()
        ids.append(i)
        meta[i] = (str(r[1]).strip() if r[1] else "", str(r[2]).strip() if r[2] else "")
        flag[i], mean[i] = norm(r[6]), norm(r[7])
    wb.close()
    return ids, meta, mean, flag


def read_physician(f):
    d = pd.read_excel(U / f)
    d = d[d["Event set"].astype(str).str.startswith("180")]
    return ({str(k): norm(v) for k, v in zip(d["Matched Event ID"], d["Clinical meaning changed?"])},
            {str(k): norm(v) for k, v in zip(d["Matched Event ID"], d["Automated flag correct?"])})


students = {"Student 1": "student_1_180.xlsx",
            "Student 2": "student_2_180.xlsx",
            "Student 3": "student_3_180.xlsx"}
S = {k: read_student(v) for k, v in students.items()}
ids, meta = S["Student 1"][0], S["Student 1"][1]
for k in S:
    assert S[k][0] == ids, f"event order differs for {k}"
R = {"Z.D. (physician)": read_physician("007b2db9-physician_2_312.xlsx"),
     "A.K.A.M. (physician)": read_physician("527a7308-physician_1_312.xlsx")}
MEAN = {k: v[0] for k, v in R.items()} | {k: v[2] for k, v in S.items()}
FLAG = {k: v[1] for k, v in R.items()} | {k: v[3] for k, v in S.items()}
names = list(MEAN)
print("ratings per reviewer:", {k: sum(1 for i in ids if MEAN[k].get(i) in LAB) for k in names})
assert all(i in MEAN["Z.D. (physician)"] for i in ids), "physician 180-set ids do not cover the student ids"

complete = [i for i in ids if all(MEAN[k].get(i) in LAB for k in names)]


def fleiss(events):
    mat = np.array([[[MEAN[k][i] for k in names].count(l) for l in LAB] for i in events], dtype=float)
    ni = mat.sum(axis=1)
    P = ((mat ** 2).sum(axis=1) - ni) / (ni * (ni - 1))
    pj = mat.sum(axis=0) / ni.sum()
    pe = (pj ** 2).sum()
    return (P.mean() - pe) / (1 - pe), P.mean()


kappa, pbar = fleiss(complete)
rng = np.random.default_rng(20260915)
boot = np.array([fleiss(list(rng.choice(complete, len(complete), replace=True)))[0] for _ in range(2000)])
lo, hi = np.percentile(boot, [2.5, 97.5])
print("Fleiss kappa %.3f (95%% CI %.2f to %.2f), agreement %.1f%%, n=%d" % (kappa, lo, hi, 100 * pbar, len(complete)))


def majority(src):
    out = {}
    for i in ids:
        c = Counter(src[k][i] for k in names if src[k].get(i) in LAB).most_common()
        out[i] = "Tied" if len(c) > 1 and c[0][1] == c[1][1] else c[0][0]
    return out


maj, fmaj = majority(MEAN), majority(FLAG)
print("majority meaning:", dict(Counter(maj.values())), "| flag majority:", dict(Counter(fmaj.values())))
for idx, lbl in ((0, "speaker"), (1, "cue")):
    g = {}
    for i in ids:
        g.setdefault(meta[i][idx], [0, 0])
        g[meta[i][idx]][1] += 1
        g[meta[i][idx]][0] += maj[i] == "Yes"
    print("  by", lbl, {a: f"{b[0]}/{b[1]}" for a, b in g.items()})

pairs = [(MEAN["Z.D. (physician)"][i], MEAN["A.K.A.M. (physician)"][i]) for i in ids]
obs = np.mean([a == b for a, b in pairs])
ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
exp = sum(ca[l] * cb[l] for l in LAB) / len(pairs) ** 2
print("physician Cohen kappa on 180: %.3f, agreement %.1f%% (%d/%d)" % ((obs - exp) / (1 - exp), 100 * obs, sum(a == b for a, b in pairs), len(pairs)))

# ---- figure 7 (house style of the original script)
for face in ("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf"):
    p = Path("/mnt/c/Windows/Fonts") / face
    if p.exists():
        font_manager.fontManager.addfont(str(p))
plt.rcParams.update({"font.family": "serif", "font.serif": ["Times New Roman", "DejaVu Serif"],
                     "mathtext.fontset": "stix", "pdf.fonttype": 42, "ps.fonttype": 42, "font.size": 9,
                     "axes.labelsize": 10, "axes.titlesize": 11, "xtick.labelsize": 9, "ytick.labelsize": 9,
                     "axes.spines.top": False, "axes.spines.right": False, "axes.edgecolor": "#9aa5ad",
                     "axes.linewidth": 0.8, "xtick.color": "#6b7780", "ytick.color": "#6b7780",
                     "axes.grid": False, "figure.dpi": 300})
fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(7.2, 3.8), gridspec_kw={"width_ratios": [1.05, 1]})
fig.subplots_adjust(left=0.16, right=0.98, bottom=0.24, top=0.85, wspace=0.42)
mat = np.array([[sum(1 for i in ids if MEAN[nm].get(i) == l) for l in LAB] for nm in names])
ax1.imshow(mat, cmap="Blues", vmin=0, vmax=mat.max() * 1.15, aspect="auto")
ax1.set_xticks(range(3), LAB)
ax1.set_yticks(range(len(names)), names)
for r in range(mat.shape[0]):
    for c in range(mat.shape[1]):
        v = mat[r, c]
        ax1.text(c, r, "{}\n({:.0f}%)".format(v, 100 * v / mat[r].sum()), ha="center", va="center", fontsize=7.8,
                 color="white" if v > mat.max() * 0.6 else "#1d262d")
ax1.set_title("a  Meaning-change ratings by reviewer", loc="left", fontweight="bold", pad=8)
for sp in ("top", "right", "left", "bottom"):
    ax1.spines[sp].set_visible(False)
ax1.tick_params(length=0)
groups = [("Patient", lambda m: m[0].upper() == "PATIENT", INK), ("Clinician", lambda m: m[0].upper() != "PATIENT", GREY),
          ("Negation", lambda m: "negation" in m[1].lower(), INK), ("Medication", lambda m: "medication" in m[1].lower(), FIRE)]
xs, pct, lab, col = [], [], [], []
for k, (nm, sel, c) in enumerate(groups):
    sub = [i for i in ids if sel(meta[i])]
    yes = sum(1 for i in sub if maj[i] == "Yes")
    xs.append(k); pct.append(100 * yes / len(sub)); col.append(c); lab.append("{}\n{}/{}".format(nm, yes, len(sub)))
bars = ax2.bar(xs, pct, width=0.62, color=col, edgecolor="#4A5A66", linewidth=0.6)
for b, v in zip(bars, pct):
    ax2.text(b.get_x() + b.get_width() / 2, v + 1.8, "{:.1f}%".format(v), ha="center", va="bottom", fontsize=8.5)
ax2.set_xticks(xs, lab)
ax2.set_ylim(0, 100)
ax2.set_ylabel("events where meaning changed (%)")
ax2.set_title("b  Majority decision by speaker and cue", loc="left", fontweight="bold", pad=8)
nconf = sum(1 for v in fmaj.values() if v == "Yes")
fig.text(0.16, 0.015, "Five reviewers, 180 flagged Fareez events; {} confirmed as genuine errors by majority.\n"
         "Fleiss' $\\kappa$ = {:.2f} (95% CI {:.2f} to {:.2f}); mean pairwise agreement {:.1f}%.".format(nconf, kappa, lo, hi, 100 * pbar),
         ha="left", va="bottom", fontsize=8.5, color="#46515c")
for ext in ("pdf", "png"):
    fig.savefig(OUT / ("fig7_clinical_agreement." + ext), bbox_inches="tight", facecolor="white")
print("wrote", OUT / "fig7_clinical_agreement.pdf", "| partial counts per reviewer:", [int(m[1]) for m in mat])

# ---- full supplementary-table statistics (flag, meaning, confidence) with updated physician ratings
CONF_LAB = ["Low", "Moderate", "High"]


def normc(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    return {"low": "Low", "moderate": "Moderate", "medium": "Moderate", "high": "High"}.get(str(v).strip().lower())


def read_student_conf(p):
    wb = openpyxl.load_workbook(D / p, data_only=True, read_only=True)
    ws = wb["ASR Comparison - 180"]
    out = {}
    for r in ws.iter_rows(min_row=2, max_col=9, values_only=True):
        if r[0] is not None:
            out[str(r[0]).strip()] = normc(r[8])
    wb.close()
    return out


def read_phys_conf(f):
    d = pd.read_excel(U / f)
    d = d[d["Event set"].astype(str).str.startswith("180")]
    return {str(k): normc(v) for k, v in zip(d["Matched Event ID"], d["Confidence"])}


CONF = {"Z.D. (physician)": read_phys_conf("007b2db9-physician_2_312.xlsx"), "A.K.A.M. (physician)": read_phys_conf("527a7308-physician_1_312.xlsx")}
CONF |= {k: read_student_conf(v) for k, v in students.items()}


def fl_generic(src, labs, events):
    mat = np.array([[[src[k][i] for k in names].count(l) for l in labs] for i in events], dtype=float)
    ni = mat.sum(axis=1)
    P = ((mat ** 2).sum(axis=1) - ni) / (ni * (ni - 1))
    pj = mat.sum(axis=0) / ni.sum()
    pe = (pj ** 2).sum()
    pe_ac1 = (pj * (1 - pj)).sum() / (len(labs) - 1)
    return (P.mean() - pe) / (1 - pe) if pe < 1 else np.nan, P.mean(), (P.mean() - pe_ac1) / (1 - pe_ac1)


print("\nSUPPLEMENTARY TABLE (five reviewers, updated physician ratings)")
for nm, src, labs in [("flag", FLAG, LAB), ("meaning", MEAN, LAB), ("confidence", CONF, CONF_LAB)]:
    ev = [i for i in ids if all(src[k].get(i) in labs for k in names)]
    k_, a_, ac1_ = fl_generic(src, labs, ev)
    rr = np.random.default_rng(20260915)
    b = [fl_generic(src, labs, list(rr.choice(ev, len(ev), replace=True)))[0] for _ in range(2000)]
    blo, bhi = np.nanpercentile(b, [2.5, 97.5])
    print(f"  {nm:10s} n={len(ev)} Fleiss={k_:.3f} (CI {blo:.2f} to {bhi:.2f}) agreement={100*a_:.1f}% AC1={ac1_:.3f}")
